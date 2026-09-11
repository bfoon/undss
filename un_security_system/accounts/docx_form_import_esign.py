"""Conservative DOCX -> editable UNPASS eSign form schema importer."""

from __future__ import annotations

import io
import re
from typing import Any, Dict, List

from . import form_pdf_esign as F

MAX_DOCX_BYTES = 12 * 1024 * 1024

_LINE_RE = re.compile(r"_{3,}|\.{4,}|-{4,}")
_FIELD_RE = re.compile(
    r"^\s*(?P<label>[^:]{2,100}?)\s*:\s*(?P<blank>_{3,}|\.{4,}|-{4,})\s*$"
)
_CHECK_SPLIT_RE = re.compile(r"[☐□☑✓✔]\s*")
_SIGNATURE_WORDS = {"signature", "signed by", "approved by", "authorised by", "authorized by"}


class WordFormImportError(ValueError):
    pass


def _uid(counter):
    return f"imp_{counter:04d}"


def _kind_for_label(label: str) -> str:
    low = label.casefold()
    if "email" in low or "e-mail" in low:
        return "email"
    if any(k in low for k in ("date", "dob", "birth")):
        return "date"
    if any(k in low for k in ("amount", "qty", "quantity", "number", "no.", "total", "price", "cost")):
        return "number"
    return "text"


def _looks_signature(text: str) -> bool:
    low = text.casefold()
    return any(word in low for word in _SIGNATURE_WORDS) and bool(_LINE_RE.search(text))


def _checkbox_options(text: str) -> List[str]:
    if not any(ch in text for ch in ("☐", "□", "☑")):
        return []
    parts = [x.strip(" :;,-\t") for x in _CHECK_SPLIT_RE.split(text)]
    return [p[:120] for p in parts if p][:50]


def _paragraph_elements(paragraph, counter: int):
    text = (paragraph.text or "").strip()
    if not text:
        return [], counter

    style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
    if style_name.casefold().startswith("heading"):
        counter += 1
        return [{
            "id": _uid(counter), "type": "heading", "width": 12,
            "text": text[:200],
            "style": {"size": "md", "color": "#005A8B", "bg": "", "align": "left"},
        }], counter

    if _looks_signature(text):
        label = re.sub(r"[_\-.]{3,}", "", text).strip(" :;-") or "Signature"
        counter += 1
        return [{
            "id": _uid(counter), "type": "signature", "width": 6,
            "label": label[:120], "role": "", "step": "", "show_date": True,
        }], counter

    options = _checkbox_options(text)
    if len(options) >= 2:
        # Text before the first checkbox is normally the question.
        first_mark = min(i for i in (text.find("☐"), text.find("□"), text.find("☑")) if i >= 0)
        label = text[:first_mark].strip(" :;-") or "Choose"
        counter += 1
        return [{
            "id": _uid(counter), "type": "checkboxes", "width": 12,
            "label": label[:200], "key": F.slug_key(label), "required": False,
            "help": "", "placeholder": "", "fill_by": "submitter", "prefill": "",
            "options": options,
        }], counter

    m = _FIELD_RE.match(text)
    if m:
        label = m.group("label").strip()
        kind = _kind_for_label(label)
        counter += 1
        el = {
            "id": _uid(counter), "type": kind, "width": 6,
            "label": label[:200], "key": F.slug_key(label), "required": False,
            "help": "", "placeholder": "", "fill_by": "submitter", "prefill": "",
        }
        if kind == "number":
            el.update({"min": None, "max": None, "decimals": 0})
        return [el], counter

    low = text.casefold()
    if ("yes" in low and "no" in low) and len(text) < 220:
        label = re.sub(r"\b(?:yes|no)\b", "", text, flags=re.I).strip(" /:;-") or text
        counter += 1
        return [{
            "id": _uid(counter), "type": "yesno", "width": 4,
            "label": label[:200], "key": F.slug_key(label), "required": False,
            "help": "", "placeholder": "", "fill_by": "submitter", "prefill": "",
        }], counter

    # A label ending with ':' is often a field prompt even without underscores.
    if text.endswith(":") and 2 <= len(text) <= 100:
        label = text[:-1].strip()
        kind = _kind_for_label(label)
        counter += 1
        el = {
            "id": _uid(counter), "type": kind, "width": 6,
            "label": label[:200], "key": F.slug_key(label), "required": False,
            "help": "", "placeholder": "", "fill_by": "submitter", "prefill": "",
        }
        if kind == "number":
            el.update({"min": None, "max": None, "decimals": 0})
        return [el], counter

    counter += 1
    return [{
        "id": _uid(counter), "type": "paragraph", "width": 12, "text": text[:4000],
        "style": {"color": "#334155", "size": "md", "bg": "", "align": "left"},
    }], counter


def _table_element(table, counter: int):
    rows = list(table.rows)
    if not rows:
        return None, counter
    max_cols = max((len(r.cells) for r in rows), default=0)
    if not max_cols:
        return None, counter

    headers = []
    for i in range(max_cols):
        text = ""
        try:
            text = (rows[0].cells[i].text or "").strip()
        except IndexError:
            pass
        headers.append(text or f"Column {i + 1}")

    cols = []
    for i, label in enumerate(headers[:10]):
        cols.append({
            "key": F.slug_key(label, f"c{i+1}"),
            "label": label[:80],
            "kind": _kind_for_label(label) if _kind_for_label(label) in {"text", "number", "date"} else "text",
            "width": 2,
        })

    counter += 1
    return {
        "id": _uid(counter), "type": "table", "width": 12,
        "label": "Items", "key": f"table_{counter}", "required": False,
        "help": "", "placeholder": "", "fill_by": "submitter", "prefill": "",
        "columns": cols, "rows": max(1, min(10, len(rows) - 1)),
        "show_total": any("amount" in c["label"].casefold() or "total" in c["label"].casefold() for c in cols),
        "allow_add": True,
    }, counter


def import_docx_form(raw: bytes, filename: str = "form.docx") -> Dict[str, Any]:
    if not raw or len(raw) > MAX_DOCX_BYTES:
        raise WordFormImportError("The Word form is empty or too large.")
    if not filename.lower().endswith(".docx"):
        raise WordFormImportError("Use a .docx Word document for editable form import.")

    try:
        from docx import Document
        doc = Document(io.BytesIO(raw))
    except Exception as exc:
        raise WordFormImportError(
            "The Word document could not be read. Confirm it is a valid .docx file."
        ) from exc

    title = ""
    try:
        title = (doc.core_properties.title or "").strip()
    except Exception:
        pass
    if not title:
        title = filename.rsplit("/", 1)[-1].rsplit(".", 1)[0].replace("_", " ").strip() or "Imported form"

    elements = []
    counter = 0

    # Preserve Word document order where possible by inspecting XML body.
    para_by_el = {p._p: p for p in doc.paragraphs}
    table_by_el = {t._tbl: t for t in doc.tables}
    for child in doc.element.body.iterchildren():
        if child in para_by_el:
            new, counter = _paragraph_elements(para_by_el[child], counter)
            elements.extend(new)
        elif child in table_by_el:
            el, counter = _table_element(table_by_el[child], counter)
            if el:
                elements.append(el)

    schema = {
        "version": 1,
        "theme": F.default_theme(),
        "header": {
            "title": title[:200],
            "subtitle": "Imported from Word — review detected fields before publishing.",
            "logo": "",
            "align": "left",
            "show_reference": True,
        },
        "elements": elements[:F.MAX_ELEMENTS],
    }
    return {
        "name": title[:150],
        "schema": F.clean_schema(schema),
        "detected_elements": len(elements),
        "warning": "Word import is heuristic. Review field types, widths, signature roles and workflow assignments.",
    }
